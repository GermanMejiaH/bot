"""Generate audit/manual_labels.xlsx for manual crop labeling.

Reads accepted candidate crop files from audit/review/accepted/ (e.g. candidate_0001_contour.png),
creates Excel table with data validation dropdowns, column auto-fitting,
and header freeze.
"""

import glob
import os
import re
from typing import Any

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

VALID_LABELS = [
    "player",
    "monster",
    "npc",
    "tree",
    "flower",
    "roof",
    "wall",
    "building",
    "bag",
    "box",
    "decoration",
    "movement_cell",
    "combat_cell",
    "unknown",
]


def main() -> None:
    accepted_dir = os.path.join("audit", "review", "accepted")
    output_excel_path = os.path.join("audit", "manual_labels.xlsx")

    png_files: list[str] = []
    if os.path.exists(accepted_dir):
        png_files = sorted(glob.glob(os.path.join(accepted_dir, "*.png")))

    if not png_files:
        crops_dir = os.path.join("audit", "crops")
        if os.path.exists(crops_dir):
            png_files = sorted(glob.glob(os.path.join(crops_dir, "*.png")))

    total_crops = len(png_files)

    rows: list[dict[str, Any]] = []
    method_counts: dict[str, int] = {"contour": 0, "hsv": 0, "combat_base": 0}
    status_counts: dict[str, int] = {"accepted": 0, "rejected": 0}

    for path in png_files:
        fn = os.path.basename(path)
        m = re.match(r"^candidate_(\d+)_(contour|hsv|combat_base)\.png$", fn)
        if m:
            cid_str, method = m.groups()
        else:
            m_cid = re.search(r"(\d+)", fn)
            cid_str = m_cid.group(1) if m_cid else "0000"
            if "hsv" in fn:
                method = "hsv"
            elif "combat_base" in fn:
                method = "combat_base"
            else:
                method = "contour"

        method_counts[method] = method_counts.get(method, 0) + 1
        status_counts["accepted"] = status_counts.get("accepted", 0) + 1

        rows.append({
            "candidate_id": cid_str,
            "filename": fn,
            "method": method,
            "accepted": "accepted",
            "label": "",
            "notes": "",
            "numerical_id": int(cid_str) if cid_str.isdigit() else 0,
        })

    # Sort rows by candidate_id numerical value, then method
    rows.sort(key=lambda r: (r["numerical_id"], r["method"]))

    # --- CREATE EXCEL WORKBOOK ---
    wb = openpyxl.Workbook()

    # 1. Labels Sheet
    ws_labels = wb.active
    ws_labels.title = "Labels"

    # 2. Metadata Sheet
    ws_meta = wb.create_sheet(title="Metadata")
    for idx, lbl in enumerate(VALID_LABELS, start=1):
        ws_meta.cell(row=idx, column=1, value=lbl)

    ws_meta.freeze_panes = "A2"

    # Populate Header
    headers = ["candidate_id", "filename", "method", "accepted", "label", "notes"]
    ws_labels.append(headers)

    # Populate Data Rows
    for r in rows:
        ws_labels.append([
            r["candidate_id"],
            r["filename"],
            r["method"],
            r["accepted"],
            r["label"],
            r["notes"],
        ])

    num_rows = len(rows) + 1  # Including header

    # Data Validation Dropdown for "label" Column (Column E)
    dv = DataValidation(
        type="list",
        formula1=f"=Metadata!$A$1:$A${len(VALID_LABELS)}",
        allow_blank=True,
    )
    dv.error = "Please select a valid label from the dropdown list."
    dv.errorTitle = "Invalid Label"
    dv.prompt = "Select label from list"
    dv.promptTitle = "Label Selection"

    ws_labels.add_data_validation(dv)
    dv.add(f"E2:E{max(2, num_rows)}")

    # Format as Excel Table
    table_ref = f"A1:F{max(2, num_rows)}"
    tab = Table(displayName="ManualLabelsTable", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    tab.tableStyleInfo = style
    ws_labels.add_table(tab)

    # Freeze Header Row
    ws_labels.freeze_panes = "A2"

    # Auto-fit Column Widths
    for col in ws_labels.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_labels.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Save Workbook
    wb.save(output_excel_path)

    print("\n=======================================================")
    print("  MANUAL LABELING SHEET GENERATION COMPLETE")
    print("=======================================================")
    print(f"Total Crops Found:    {total_crops}")
    print(f"Generated File Path:  {output_excel_path}")
    print("-------------------------------------------------------")
    print("Distribution by Method:")
    for m_name, m_cnt in method_counts.items():
        pct = round(m_cnt / float(max(1, total_crops)) * 100.0, 1)
        print(f"  - {m_name:<12}: {m_cnt:>4} ({pct}%)")
    print("-------------------------------------------------------")
    print("Distribution by Status:")
    for s_name, s_cnt in status_counts.items():
        pct = round(s_cnt / float(max(1, total_crops)) * 100.0, 1)
        print(f"  - {s_name:<12}: {s_cnt:>4} ({pct}%)")
    print("=======================================================\n")


if __name__ == "__main__":
    main()
