"""Phase 4D.1C — Asset Generator for Manual Positive Audit Preparation.

Creates:
1. audit/manual_positive_review_v2.xlsx
   Columns: sample_id, probability, method, bbox_width, bbox_height, review_label
   Dropdown: VALID_ENTITY, PARTIAL_ENTITY, DECORATION, OVERSIZED_BBOX, UNCERTAIN

2. audit/positive_audit_statistics.py
   Script to analyze completed manual_positive_review_v2.xlsx and produce:
   - audit/positive_audit_summary.json
   - audit/positive_audit_summary.md

# ruff: noqa: N806, N803
"""

import csv
import glob
import os
from typing import Any

import cv2
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from dta.config.settings import get_settings
from dta.detectors.character_detector import CharacterDetector
from dta.perception.visual_classifier import VisualCandidateClassifier
from dta.vision.map_roi_extractor import MapROIExtractor

ALLOWED_REVIEW_LABELS = [
    "VALID_ENTITY",
    "PARTIAL_ENTITY",
    "DECORATION",
    "OVERSIZED_BBOX",
    "UNCERTAIN",
]


def collect_top100_positives() -> list[dict[str, Any]]:
    """Collect top 100 highest probability candidates."""
    csv_path = os.path.join("audit", "top100_positive_audit.csv")

    # If top100_positive_audit.csv exists, load directly from CSV
    if os.path.exists(csv_path):
        items: list[dict[str, Any]] = []
        with open(csv_path, encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                items.append({
                    "sample_id": row["sample_id"],
                    "probability": float(row["probability"]),
                    "method": row["method"],
                    "bbox_width": int(row["bbox_width"]),
                    "bbox_height": int(row["bbox_height"]),
                })
        if len(items) == 100:
            return items

    # Fallback to runtime extraction
    settings = get_settings()
    detector = CharacterDetector()
    roi_extractor = MapROIExtractor(settings=settings)
    classifier = VisualCandidateClassifier(model_path="models/entity_classifier.pkl", threshold=0.50)

    exp_folder = os.path.join("dataset", "exploration")
    com_folder = os.path.join("dataset", "combat")

    all_candidates: list[dict[str, Any]] = []
    cand_counter = 1

    folders = [("exploration", exp_folder), ("combat", com_folder)]
    for _domain, folder in folders:
        png_files = sorted(glob.glob(os.path.join(folder, "*.png")))
        for fpath in png_files:
            frame = cv2.imread(fpath)
            if frame is None or frame.size == 0:
                continue

            roi_frame = roi_extractor.apply_roi_mask(frame)
            trace = detector.detect_characters_with_trace(roi_frame)
            raw_characters = trace.get("raw_characters", [])

            for det in raw_characters:
                bbox = det.bbox
                x, y, w, h = max(0, bbox.x), max(0, bbox.y), bbox.w, bbox.h
                img_h, img_w = frame.shape[:2]
                x2, y2 = min(img_w, x + w), min(img_h, y + h)

                crop = frame[y:y2, x:x2]
                if crop.size == 0:
                    continue

                proba = float(classifier.predict_proba(crop))
                all_candidates.append({
                    "sample_id": f"SAMP_{cand_counter:04d}",
                    "probability": proba,
                    "method": det.method,
                    "bbox_width": w,
                    "bbox_height": h,
                })
                cand_counter += 1

    positives = [c for c in all_candidates if c["probability"] >= 0.50]
    positives.sort(key=lambda x: x["probability"], reverse=True)
    return positives[:100]


def create_manual_positive_review_v2_excel(items: list[dict[str, Any]], output_path: str) -> None:
    """Generate audit/manual_positive_review_v2.xlsx with data validation dropdowns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Positive Review v2"

    headers = [
        "sample_id",
        "probability",
        "method",
        "bbox_width",
        "bbox_height",
        "review_label",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 7):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Setup Data Validation for review_label dropdown
    label_formula = '"' + ",".join(ALLOWED_REVIEW_LABELS) + '"'
    dv = DataValidation(type="list", formula1=label_formula, allow_blank=True)
    ws.add_data_validation(dv)

    for item in items:
        ws.append([
            item["sample_id"],
            round(item["probability"], 4),
            item["method"],
            item["bbox_width"],
            item["bbox_height"],
            "",  # review_label initially blank
        ])
        current_row = ws.max_row

        ws.cell(row=current_row, column=1).alignment = align_left
        ws.cell(row=current_row, column=2).alignment = align_right
        ws.cell(row=current_row, column=2).number_format = "0.0000"
        ws.cell(row=current_row, column=3).alignment = align_center
        ws.cell(row=current_row, column=4).alignment = align_right
        ws.cell(row=current_row, column=5).alignment = align_right
        ws.cell(row=current_row, column=6).alignment = align_center

    row_end = ws.max_row
    if row_end >= 2:
        dv.add(f"F2:F{row_end}")

    # Set column dimensions
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 24

    wb.save(output_path)
    print(f"Generated Excel review sheet at '{output_path}'.")


def create_positive_audit_statistics_script(script_path: str) -> None:
    """Generate audit/positive_audit_statistics.py script."""
    script_code = r'''"""Phase 4D.1C — Positive Audit Statistics Analyzer.

Reads audit/manual_positive_review_v2.xlsx, computes category frequencies,
and exports:
- audit/positive_audit_summary.json
- audit/positive_audit_summary.md

# ruff: noqa: N806, N803
"""

import json
import os

import openpyxl

ALLOWED_CATEGORIES = [
    "VALID_ENTITY",
    "PARTIAL_ENTITY",
    "DECORATION",
    "OVERSIZED_BBOX",
    "UNCERTAIN",
]


def analyze_positive_audit(excel_path: str) -> dict:
    """Analyze annotations in manual_positive_review_v2.xlsx."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at '{excel_path}'")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    total_samples = 0
    annotated_samples = 0
    unannotated_samples = 0

    category_counts = {cat: 0 for cat in ALLOWED_CATEGORIES}
    category_counts["OTHER_UNCLASSIFIED"] = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        total_samples += 1
        review_label = str(row[5]).strip() if len(row) >= 6 and row[5] is not None else ""

        if review_label == "" or review_label.lower() == "none":
            unannotated_samples += 1
        elif review_label in category_counts:
            category_counts[review_label] += 1
            annotated_samples += 1
        else:
            category_counts["OTHER_UNCLASSIFIED"] += 1
            annotated_samples += 1

    denom = max(1, annotated_samples)
    category_percentages = {
        cat: round((count / float(denom)) * 100.0, 2)
        for cat, count in category_counts.items()
    }

    # Determine dominant error source (excluding VALID_ENTITY and UNCERTAIN)
    error_cats = {k: v for k, v in category_counts.items() if k not in {"VALID_ENTITY", "UNCERTAIN"}}
    dominant_error = max(error_cats.items(), key=lambda x: x[1])[0] if error_cats and max(error_cats.values()) > 0 else "NONE"

    summary_data = {
        "total_samples": total_samples,
        "annotated_samples": annotated_samples,
        "unannotated_samples": unannotated_samples,
        "category_counts": category_counts,
        "category_percentages": category_percentages,
        "dominant_error_source": dominant_error,
    }
    return summary_data


def export_summary_json(data: dict, json_path: str) -> None:
    """Export summary statistics to JSON."""
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print(f"Exported audit statistics JSON to '{json_path}'.")


def export_summary_md(data: dict, md_path: str) -> None:
    """Export audit statistics report to Markdown."""
    total = data["total_samples"]
    annotated = data["annotated_samples"]
    unannotated = data["unannotated_samples"]
    counts = data["category_counts"]
    pcts = data["category_percentages"]
    dominant = data["dominant_error_source"]

    md = "# Phase 4D.1C — Positive Candidate Contamination Audit Summary\n\n"
    md += "## Executive Summary\n\n"
    md += f"A total of **{annotated} / {total} accepted positive predictions** ($P \\ge 0.50$) were annotated and categorized to measure dataset contamination sources.\n\n"
    md += f"- **Annotated Samples**: `{annotated}`\n"
    md += f"- **Unannotated Samples**: `{unannotated}`\n"
    md += f"- **Dominant Contamination Source**: `{dominant}`\n\n"

    md += "## 1. Category Frequency Breakdown\n\n"
    md += "| Review Category | Count | Percentage Share (%) | Visual Impact & Interpretation |\n"
    md += "| :--- | :---: | :---: | :--- |\n"
    md += f"| **`VALID_ENTITY`** | `{counts.get('VALID_ENTITY', 0)}` | `{pcts.get('VALID_ENTITY', 0.0)}%` | True valid player character, monster, or NPC sprite core. |\n"
    md += f"| **`PARTIAL_ENTITY`** | `{counts.get('PARTIAL_ENTITY', 0)}` | `{pcts.get('PARTIAL_ENTITY', 0.0)}%` | Character crop with low target occupancy (<30% body core). |\n"
    md += f"| **`DECORATION`** | `{counts.get('DECORATION', 0)}` | `{pcts.get('DECORATION', 0.0)}%` | Scenery prop, statue, lamp post, wall column, or plant flora. |\n"
    md += f"| **`OVERSIZED_BBOX`** | `{counts.get('OVERSIZED_BBOX', 0)}` | `{pcts.get('OVERSIZED_BBOX', 0.0)}%` | Bounding box covering multiple ground cells or terrain tiles. |\n"
    md += f"| **`UNCERTAIN`** | `{counts.get('UNCERTAIN', 0)}` | `{pcts.get('UNCERTAIN', 0.0)}%` | Ambiguous crop requiring secondary visual context. |\n"
    md += f"| **Total Annotated** | `{annotated}` | `100.0%` | Audited positive candidates. |\n\n"

    md += "## 2. Root Cause Analysis & Primary Contamination Drivers\n\n"
    if dominant == "DECORATION":
        md += "The audit indicates that **`DECORATION` contamination** is the primary driver of false positive acceptances. High-contrast vertical scenery objects (statues, lamp posts, wall columns) produce visual feature embeddings that confuse the linear classifier.\n"
    elif dominant == "PARTIAL_ENTITY":
        md += "The audit indicates that **`PARTIAL_ENTITY` crops** dominate false positive acceptances. Bounding box candidate extraction often captures ground floor tiles with only a tiny fraction of an entity edge.\n"
    elif dominant == "OVERSIZED_BBOX":
        md += "The audit indicates that **`OVERSIZED_BBOX` candidates** dominate contamination. Large multi-tile bounding boxes encompassing floor terrain and multiple objects are being accepted as single entities.\n"
    else:
        md += "Contamination is distributed across multiple categories including decorations, partial entity crops, and oversized bounding boxes.\n"

    md += "\n## 3. Recommended Remediation Strategy\n\n"
    md += "1. **Deterministic BBox Pre-Filtering**: Add size floor (`area >= 400`), size ceiling (`area <= 12000`), and ROI floor limits in `CharacterDetector` to purge oversized multi-tiles and micro-glints prior to visual classification.\n"
    md += "2. **Targeted Negative Retraining**: Expand negative dataset with explicit `DECORATION` and `PARTIAL_ENTITY` crops to refine classifier decision boundaries in Phase 4D.2.\n"

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md)
    print(f"Exported audit statistics Markdown summary to '{md_path}'.")


def main() -> None:
    excel_path = os.path.join("audit", "manual_positive_review_v2.xlsx")
    json_path = os.path.join("audit", "positive_audit_summary.json")
    md_path = os.path.join("audit", "positive_audit_summary.md")

    print("\n=======================================================")
    print("  PHASE 4D.1C POSITIVE AUDIT STATISTICS ANALYZER")
    print("=======================================================")

    summary_data = analyze_positive_audit(excel_path)
    export_summary_json(summary_data, json_path)
    export_summary_md(summary_data, md_path)

    print("\nPhase 4D.1C Statistics Analysis Successfully Completed!")


if __name__ == "__main__":
    main()
'''

    with open(script_path, "w", encoding="utf-8") as f:
        f.write(script_code)
    print(f"Generated audit statistics script at '{script_path}'.")


def main() -> None:
    """Execute Phase 4D.1C preparation task."""
    print("\n=======================================================")
    print("  PHASE 4D.1C MANUAL POSITIVE AUDIT PREPARATION")
    print("=======================================================")

    os.makedirs("audit", exist_ok=True)

    items = collect_top100_positives()
    print(f"Collected top 100 positive candidate items: {len(items)} samples.")

    # 1. Generate audit/manual_positive_review_v2.xlsx
    excel_v2_path = os.path.join("audit", "manual_positive_review_v2.xlsx")
    create_manual_positive_review_v2_excel(items, excel_v2_path)

    # 2. Generate audit/positive_audit_statistics.py
    stats_script_path = os.path.join("audit", "positive_audit_statistics.py")
    create_positive_audit_statistics_script(stats_script_path)

    print("\n=======================================================")
    print("  PHASE 4D.1C AUDIT ASSETS PREPARED SUCCESSFULLY!")
    print("=======================================================")


if __name__ == "__main__":
    main()
