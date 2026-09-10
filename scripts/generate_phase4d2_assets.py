"""Phase 4D.2 — Bounding Box Failure Taxonomy Audit Generator.

Processes manual_positive_review_v2.xlsx annotations (PARTIAL_ENTITY & OVERSIZED_BBOX items):
1. Generates audit/partial_entity_grid.png
2. Generates audit/oversized_bbox_grid.png
3. Generates audit/bbox_failure_review.xlsx
4. Generates audit/bbox_failure_statistics.py

# ruff: noqa: N806, N803
"""

import glob
import os
from typing import Any

import cv2
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from dta.config.settings import get_settings
from dta.detectors.character_detector import CharacterDetector
from dta.perception.visual_classifier import VisualCandidateClassifier
from dta.vision.map_roi_extractor import MapROIExtractor

ALLOWED_FAILURE_TYPES = [
    "LEFT_CLIPPED",
    "RIGHT_CLIPPED",
    "TOP_CLIPPED",
    "BOTTOM_CLIPPED",
    "PET_ONLY",
    "MULTI_ENTITY_FRAGMENT",
    "MULTI_TILE",
    "GROUND_TILE",
    "ENTITY_PLUS_TERRAIN",
    "ENTITY_PLUS_UI",
    "UNCERTAIN",
]


def collect_runtime_candidates_map() -> dict[str, dict[str, Any]]:
    """Collect runtime candidates mapped by sample_id."""
    settings = get_settings()
    detector = CharacterDetector()
    roi_extractor = MapROIExtractor(settings=settings)
    classifier = VisualCandidateClassifier(model_path="models/entity_classifier.pkl", threshold=0.50)

    exp_folder = os.path.join("dataset", "exploration")
    com_folder = os.path.join("dataset", "combat")

    cand_map: dict[str, dict[str, Any]] = {}
    cand_counter = 1

    folders = [("exploration", exp_folder), ("combat", com_folder)]
    for _domain, folder in folders:
        png_files = sorted(glob.glob(os.path.join(folder, "*.png")))
        for fpath in png_files:
            frame = cv2.imread(fpath)
            if frame is None or frame.size == 0:
                continue

            roi_frame = roi_extractor.apply_roi_mask(frame)
            trace = detector.detect_characters_with_trace(roi_frame)
            raw_characters = trace.get("raw_characters", [])

            for det in raw_characters:
                bbox = det.bbox
                x, y, w, h = max(0, bbox.x), max(0, bbox.y), bbox.w, bbox.h
                img_h, img_w = frame.shape[:2]
                x2, y2 = min(img_w, x + w), min(img_h, y + h)

                crop = frame[y:y2, x:x2]
                if crop.size == 0:
                    continue

                proba = float(classifier.predict_proba(crop))
                sample_id = f"SAMP_{cand_counter:04d}"
                cand_map[sample_id] = {
                    "sample_id": sample_id,
                    "probability": proba,
                    "method": det.method,
                    "bbox_width": w,
                    "bbox_height": h,
                    "crop": crop,
                }
                cand_counter += 1

    return cand_map


def load_review_v2_annotations(excel_path: str) -> list[dict[str, Any]]:
    """Read manual_positive_review_v2.xlsx and return annotated candidate records."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Missing review spreadsheet at '{excel_path}'")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    records: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        sample_id = str(row[0]).strip()
        probability = float(row[1]) if row[1] is not None else 0.0
        method = str(row[2]).strip() if len(row) >= 3 and row[2] else ""
        bbox_w = int(row[3]) if len(row) >= 4 and row[3] else 0
        bbox_h = int(row[4]) if len(row) >= 5 and row[4] else 0
        review_label = str(row[5]).strip() if len(row) >= 6 and row[5] else ""

        records.append({
            "sample_id": sample_id,
            "probability": probability,
            "method": method,
            "bbox_width": bbox_w,
            "bbox_height": bbox_h,
            "review_label": review_label,
        })

    return records


def render_image_grid(
    items: list[dict[str, Any]],
    title: str,
    output_path: str,
    cols: int = 5,
    cell_w: int = 180,
    cell_h: int = 220,
) -> None:
    """Render annotated image grid for failure candidates."""
    num_items = len(items)
    if num_items == 0:
        print(f"Warning: No items to render for grid '{output_path}'.")
        return

    rows = int(np.ceil(num_items / float(cols)))
    margin = 30
    header_h = 70

    grid_w = margin * 2 + cols * cell_w + (cols - 1) * 8
    grid_h = header_h + margin * 2 + rows * cell_h + (rows - 1) * 8

    canvas = np.full((grid_h, grid_w, 3), (24, 26, 32), dtype=np.uint8)

    # Header bar
    cv2.rectangle(canvas, (0, 0), (grid_w, header_h), (36, 40, 50), -1)
    cv2.putText(
        canvas,
        title,
        (margin, 44),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.80,
        (240, 240, 240),
        2,
        cv2.LINE_AA,
    )
    cv2.putText(
        canvas,
        f"Count: {num_items}",
        (grid_w - margin - 150, 44),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.60,
        (180, 200, 220),
        1,
        cv2.LINE_AA,
    )

    for idx, item in enumerate(items):
        r = idx // cols
        c = idx % cols

        x_start = margin + c * (cell_w + 8)
        y_start = header_h + margin + r * (cell_h + 8)
        x_end = x_start + cell_w
        y_end = y_start + cell_h

        # Cell background card
        cv2.rectangle(canvas, (x_start, y_start), (x_end, y_end), (40, 44, 54), -1)
        cv2.rectangle(canvas, (x_start, y_start), (x_end, y_end), (65, 72, 88), 1)

        # Crop area
        crop_box_w = cell_w - 12
        crop_box_h = 130
        crop_img = item.get("crop")

        if crop_img is not None and crop_img.size > 0:
            ch, cw = crop_img.shape[:2]
            scale = min(crop_box_w / float(cw), crop_box_h / float(ch))
            nw = max(1, int(cw * scale))
            nh = max(1, int(ch * scale))

            resized_crop = cv2.resize(crop_img, (nw, nh), interpolation=cv2.INTER_AREA)

            off_x = x_start + 6 + (crop_box_w - nw) // 2
            off_y = y_start + 6 + (crop_box_h - nh) // 2

            canvas[off_y : off_y + nh, off_x : off_x + nw] = resized_crop
            cv2.rectangle(
                canvas,
                (x_start + 5, y_start + 5),
                (x_start + 5 + crop_box_w, y_start + 5 + crop_box_h),
                (80, 90, 110),
                1,
            )

        # Metadata text
        text_y = y_start + 152
        cv2.putText(
            canvas,
            item["sample_id"],
            (x_start + 8, text_y),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.44,
            (255, 215, 0),
            1,
            cv2.LINE_AA,
        )

        proba_val = item["probability"]
        cv2.putText(
            canvas,
            f"P: {proba_val:.4f}",
            (x_start + 8, text_y + 18),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.42,
            (100, 230, 100),
            1,
            cv2.LINE_AA,
        )

        size_str = f"{item['bbox_width']}x{item['bbox_height']}"
        cv2.putText(
            canvas,
            f"Size: {size_str}",
            (x_start + 8, text_y + 34),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.38,
            (200, 200, 200),
            1,
            cv2.LINE_AA,
        )

        cv2.putText(
            canvas,
            f"Det: {item['method']}",
            (x_start + 8, text_y + 48),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.36,
            (170, 180, 195),
            1,
            cv2.LINE_AA,
        )

    cv2.imwrite(output_path, canvas)
    print(f"Saved visual grid to '{output_path}'.")


def create_bbox_failure_review_excel(items: list[dict[str, Any]], output_path: str) -> None:
    """Generate audit/bbox_failure_review.xlsx with Data Validation dropdowns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BBox Failure Review"

    headers = [
        "sample_id",
        "probability",
        "method",
        "bbox_width",
        "bbox_height",
        "failure_type",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 7):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Setup Data Validation for failure_type dropdown
    fail_formula = '"' + ",".join(ALLOWED_FAILURE_TYPES) + '"'
    dv = DataValidation(type="list", formula1=fail_formula, allow_blank=True)
    ws.add_data_validation(dv)

    for item in items:
        ws.append([
            item["sample_id"],
            round(item["probability"], 4),
            item["method"],
            item["bbox_width"],
            item["bbox_height"],
            "",  # failure_type initially blank
        ])
        current_row = ws.max_row

        ws.cell(row=current_row, column=1).alignment = align_left
        ws.cell(row=current_row, column=2).alignment = align_right
        ws.cell(row=current_row, column=2).number_format = "0.0000"
        ws.cell(row=current_row, column=3).alignment = align_center
        ws.cell(row=current_row, column=4).alignment = align_right
        ws.cell(row=current_row, column=5).alignment = align_right
        ws.cell(row=current_row, column=6).alignment = align_center

    row_end = ws.max_row
    if row_end >= 2:
        dv.add(f"F2:F{row_end}")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 26

    wb.save(output_path)
    print(f"Generated Excel review sheet at '{output_path}'.")


def create_bbox_failure_statistics_script(script_path: str) -> None:
    """Generate audit/bbox_failure_statistics.py script."""
    script_code = r'''"""Phase 4D.2 — Bounding Box Failure Statistics Analyzer.

Reads audit/bbox_failure_review.xlsx, summarizes failure distributions,
and exports:
- audit/bbox_failure_summary.json
- audit/bbox_failure_summary.md

# ruff: noqa: N806, N803
"""

import json
import os

import openpyxl

ALLOWED_FAILURE_TYPES = [
    "LEFT_CLIPPED",
    "RIGHT_CLIPPED",
    "TOP_CLIPPED",
    "BOTTOM_CLIPPED",
    "PET_ONLY",
    "MULTI_ENTITY_FRAGMENT",
    "MULTI_TILE",
    "GROUND_TILE",
    "ENTITY_PLUS_TERRAIN",
    "ENTITY_PLUS_UI",
    "UNCERTAIN",
]


def analyze_bbox_failures(excel_path: str) -> dict:
    """Analyze annotations in bbox_failure_review.xlsx."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at '{excel_path}'")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    total_samples = 0
    annotated_samples = 0
    unannotated_samples = 0

    failure_counts = {ft: 0 for ft in ALLOWED_FAILURE_TYPES}
    failure_counts["OTHER_UNCLASSIFIED"] = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        total_samples += 1
        failure_type = str(row[5]).strip() if len(row) >= 6 and row[5] is not None else ""

        if failure_type == "" or failure_type.lower() == "none":
            unannotated_samples += 1
        elif failure_type in failure_counts:
            failure_counts[failure_type] += 1
            annotated_samples += 1
        else:
            failure_counts["OTHER_UNCLASSIFIED"] += 1
            annotated_samples += 1

    denom = max(1, annotated_samples)
    failure_percentages = {
        ft: round((count / float(denom)) * 100.0, 2)
        for ft, count in failure_counts.items()
    }

    # Dominant failure mode
    non_uncertain = {k: v for k, v in failure_counts.items() if k != "UNCERTAIN"}
    dominant = max(non_uncertain.items(), key=lambda x: x[1])[0] if non_uncertain and max(non_uncertain.values()) > 0 else "NONE"

    return {
        "total_samples": total_samples,
        "annotated_samples": annotated_samples,
        "unannotated_samples": unannotated_samples,
        "failure_counts": failure_counts,
        "failure_percentages": failure_percentages,
        "dominant_failure_mode": dominant,
    }


def export_summary_json(data: dict, json_path: str) -> None:
    """Export failure summary to JSON."""
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print(f"Exported BBox failure summary JSON to '{json_path}'.")


def export_summary_md(data: dict, md_path: str) -> None:
    """Export failure summary report to Markdown."""
    total = data["total_samples"]
    annotated = data["annotated_samples"]
    unannotated = data["unannotated_samples"]
    counts = data["failure_counts"]
    pcts = data["failure_percentages"]
    dominant = data["dominant_failure_mode"]

    md = "# Phase 4D.2 — Bounding Box Geometric Failure Taxonomy Audit Summary\n\n"
    md += "## Executive Summary\n\n"
    md += f"A detailed geometric failure mode audit was performed across **{annotated} / {total} flagged candidate bounding boxes** (`PARTIAL_ENTITY` and `OVERSIZED_BBOX`).\n\n"
    md += f"- **Total Failure Candidates Audited**: `{total}`\n"
    md += f"- **Annotated Samples**: `{annotated}`\n"
    md += f"- **Unannotated Samples**: `{unannotated}`\n"
    md += f"- **Dominant Geometric Failure Mode**: `{dominant}`\n\n"

    md += "## 1. Geometric Failure Mode Distribution\n\n"
    md += "| Failure Mode | Count | Percentage Share (%) | Visual Manifestation & Cause |\n"
    md += "| :--- | :---: | :---: | :--- |\n"
    for ft in ALLOWED_FAILURE_TYPES:
        count = counts.get(ft, 0)
        pct = pcts.get(ft, 0.0)
        md += f"| **`{ft}`** | `{count}` | `{pct}%` | Geometric candidate bound distortion. |\n"
    md += f"| **Total Annotated** | `{annotated}` | `100.0%` | Audited geometry candidates. |\n\n"

    md += "## 2. Root Cause Analysis & Detector Recommendations\n\n"
    md += f"The dominant geometric failure mode identified is **`{dominant}`**.\n\n"
    md += "### Key Findings & Detector Refinement Roadmap:\n"
    md += "1. **Boundary Clipping Mitigation**: If clipping (`TOP_CLIPPED`, `BOTTOM_CLIPPED`) dominates, adjust contour expansion padding and HSV bounding box expansion in `CharacterDetector`.\n"
    md += "2. **Multi-Tile Terrain Purging**: If `MULTI_TILE` or `ENTITY_PLUS_TERRAIN` dominates, enforce bounding box area ceiling (`area <= 12,000`) and aspect ratio caps before visual classification.\n"
    md += "3. **Pet/Fragment Filtering**: If `PET_ONLY` or `MULTI_ENTITY_FRAGMENT` dominates, refine color thresholds and contour hierarchy checks.\n"

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md)
    print(f"Exported BBox failure summary Markdown report to '{md_path}'.")


def main() -> None:
    excel_path = os.path.join("audit", "bbox_failure_review.xlsx")
    json_path = os.path.join("audit", "bbox_failure_summary.json")
    md_path = os.path.join("audit", "bbox_failure_summary.md")

    print("\n=======================================================")
    print("  PHASE 4D.2 BBOX FAILURE STATISTICS ANALYZER")
    print("=======================================================")

    summary_data = analyze_bbox_failures(excel_path)
    export_summary_json(summary_data, json_path)
    export_summary_md(summary_data, md_path)

    print("\nPhase 4D.2 BBox Failure Analysis Successfully Completed!")


if __name__ == "__main__":
    main()
'''

    with open(script_path, "w", encoding="utf-8") as f:
        f.write(script_code)
    print(f"Generated failure statistics analyzer script at '{script_path}'.")


def main() -> None:
    """Execute Phase 4D.2 Bounding Box Failure Taxonomy Audit Generator."""
    print("\n=======================================================")
    print("  PHASE 4D.2 BBOX FAILURE TAXONOMY AUDIT GENERATOR")
    print("=======================================================")

    os.makedirs("audit", exist_ok=True)

    # 1. Load manual_positive_review_v2.xlsx
    excel_v2_path = os.path.join("audit", "manual_positive_review_v2.xlsx")
    review_records = load_review_v2_annotations(excel_v2_path)

    # Filter PARTIAL_ENTITY and OVERSIZED_BBOX
    partial_items = [r for r in review_records if r["review_label"] == "PARTIAL_ENTITY"]
    oversized_items = [r for r in review_records if r["review_label"] == "OVERSIZED_BBOX"]

    print(f"Extracted {len(partial_items)} PARTIAL_ENTITY samples.")
    print(f"Extracted {len(oversized_items)} OVERSIZED_BBOX samples.")

    # Load runtime crop images map
    cand_map = collect_runtime_candidates_map()

    # Attach crops to items
    for item in partial_items:
        sid = item["sample_id"]
        if sid in cand_map:
            item["crop"] = cand_map[sid]["crop"]

    for item in oversized_items:
        sid = item["sample_id"]
        if sid in cand_map:
            item["crop"] = cand_map[sid]["crop"]

    # 2. Generate Image Grids
    partial_grid_path = os.path.join("audit", "partial_entity_grid.png")
    render_image_grid(
        items=partial_items,
        title="Phase 4D.2 — PARTIAL_ENTITY Candidates (23 Samples)",
        output_path=partial_grid_path,
        cols=5,
    )

    oversized_grid_path = os.path.join("audit", "oversized_bbox_grid.png")
    render_image_grid(
        items=oversized_items,
        title="Phase 4D.2 — OVERSIZED_BBOX Candidates (16 Samples)",
        output_path=oversized_grid_path,
        cols=4,
    )

    # 3. Create audit/bbox_failure_review.xlsx
    failure_review_items = partial_items + oversized_items
    failure_review_items.sort(key=lambda x: x["sample_id"])
    excel_fail_path = os.path.join("audit", "bbox_failure_review.xlsx")
    create_bbox_failure_review_excel(failure_review_items, excel_fail_path)

    # 4. Generate audit/bbox_failure_statistics.py
    stats_script_path = os.path.join("audit", "bbox_failure_statistics.py")
    create_bbox_failure_statistics_script(stats_script_path)

    print("\n=======================================================")
    print("  PHASE 4D.2 AUDIT ASSETS GENERATED SUCCESSFULLY!")
    print("=======================================================")


if __name__ == "__main__":
    main()
