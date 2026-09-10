"""Phase 4D.1B — Manual Contamination Verification Asset Generator.

Generates visual review grids and Excel manual review worksheets for:
1. Top 100 accepted positive predictions (P >= 0.50) -> audit/top100_positive_grid.png & audit/manual_positive_review.xlsx
2. All borderline predictions (0.40 <= P <= 0.60) -> audit/borderline_review_grid.png & audit/manual_borderline_review.xlsx

# ruff: noqa: N806, N803
"""

import glob
import os
from typing import Any

import cv2
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from dta.config.settings import get_settings
from dta.detectors.character_detector import CharacterDetector
from dta.perception.visual_classifier import VisualCandidateClassifier
from dta.vision.map_roi_extractor import MapROIExtractor

ALLOWED_CATEGORIES = [
    "valid_entity",
    "decoration",
    "plant",
    "ui",
    "effect",
    "terrain",
    "oversized_crop",
    "uncertain",
]


def collect_runtime_candidates(
    exploration_folder: str,
    combat_folder: str,
    classifier: VisualCandidateClassifier,
    detector: CharacterDetector,
    roi_extractor: MapROIExtractor,
) -> list[dict[str, Any]]:
    """Evaluate candidates across exploration and combat screenshot folders."""
    all_candidates: list[dict[str, Any]] = []
    cand_counter = 1

    folders = [("exploration", exploration_folder), ("combat", combat_folder)]

    for domain, folder in folders:
        png_files = sorted(glob.glob(os.path.join(folder, "*.png")))
        for fpath in png_files:
            fn = os.path.basename(fpath)
            frame_id = os.path.splitext(fn)[0]
            frame = cv2.imread(fpath)
            if frame is None or frame.size == 0:
                continue

            roi_frame = roi_extractor.apply_roi_mask(frame)
            trace = detector.detect_characters_with_trace(roi_frame)
            raw_characters = trace.get("raw_characters", [])

            for det in raw_characters:
                bbox = det.bbox
                x, y, w, h = max(0, bbox.x), max(0, bbox.y), bbox.w, bbox.h
                img_h, img_w = frame.shape[:2]
                x2, y2 = min(img_w, x + w), min(img_h, y + h)

                crop = frame[y:y2, x:x2]
                if crop.size == 0:
                    continue

                proba = float(classifier.predict_proba(crop))
                area = w * h
                aspect = float(w / float(max(1, h)))

                all_candidates.append({
                    "sample_id": f"SAMP_{cand_counter:04d}",
                    "domain": domain,
                    "frame_id": frame_id,
                    "filename": fn,
                    "method": det.method,
                    "bbox": [x, y, w, h],
                    "bbox_width": w,
                    "bbox_height": h,
                    "area": area,
                    "aspect_ratio": aspect,
                    "probability": proba,
                    "crop": crop,
                })
                cand_counter += 1

    return all_candidates


def render_image_grid(
    items: list[dict[str, Any]],
    title: str,
    output_path: str,
    cols: int = 10,
    cell_w: int = 180,
    cell_h: int = 220,
) -> None:
    """Render a clean grid image with annotated crops and metadata."""
    num_items = len(items)
    if num_items == 0:
        print(f"Warning: No items to render for grid '{output_path}'.")
        return

    rows = int(np.ceil(num_items / float(cols)))
    margin = 30
    header_h = 70

    grid_w = margin * 2 + cols * cell_w + (cols - 1) * 8
    grid_h = header_h + margin * 2 + rows * cell_h + (rows - 1) * 8

    # Create dark charcoal canvas
    canvas = np.full((grid_h, grid_w, 3), (24, 26, 32), dtype=np.uint8)

    # Header bar
    cv2.rectangle(canvas, (0, 0), (grid_w, header_h), (36, 40, 50), -1)
    cv2.putText(
        canvas,
        title,
        (margin, 44),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.85,
        (240, 240, 240),
        2,
        cv2.LINE_AA,
    )
    cv2.putText(
        canvas,
        f"Total Samples: {num_items}",
        (grid_w - margin - 220, 44),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.60,
        (180, 200, 220),
        1,
        cv2.LINE_AA,
    )

    for idx, item in enumerate(items):
        r = idx // cols
        c = idx % cols

        x_start = margin + c * (cell_w + 8)
        y_start = header_h + margin + r * (cell_h + 8)
        x_end = x_start + cell_w
        y_end = y_start + cell_h

        # Cell background card
        cv2.rectangle(canvas, (x_start, y_start), (x_end, y_end), (40, 44, 54), -1)
        cv2.rectangle(canvas, (x_start, y_start), (x_end, y_end), (65, 72, 88), 1)

        # Crop image area (top 135 px of cell)
        crop_box_w = cell_w - 12
        crop_box_h = 130
        crop_img = item["crop"]

        if crop_img is not None and crop_img.size > 0:
            ch, cw = crop_img.shape[:2]
            scale = min(crop_box_w / float(cw), crop_box_h / float(ch))
            nw = max(1, int(cw * scale))
            nh = max(1, int(ch * scale))

            resized_crop = cv2.resize(crop_img, (nw, nh), interpolation=cv2.INTER_AREA)

            # Center crop inside crop box
            off_x = x_start + 6 + (crop_box_w - nw) // 2
            off_y = y_start + 6 + (crop_box_h - nh) // 2

            canvas[off_y : off_y + nh, off_x : off_x + nw] = resized_crop
            cv2.rectangle(
                canvas,
                (x_start + 5, y_start + 5),
                (x_start + 5 + crop_box_w, y_start + 5 + crop_box_h),
                (80, 90, 110),
                1,
            )

        # Metadata text area below crop (bottom 80 px)
        text_y = y_start + 152
        cv2.putText(
            canvas,
            item["sample_id"],
            (x_start + 8, text_y),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.44,
            (255, 215, 0),  # Gold color for Sample ID
            1,
            cv2.LINE_AA,
        )

        proba_val = item["probability"]
        proba_color = (100, 230, 100) if proba_val >= 0.50 else (100, 180, 255)
        cv2.putText(
            canvas,
            f"P: {proba_val:.4f}",
            (x_start + 8, text_y + 18),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.42,
            proba_color,
            1,
            cv2.LINE_AA,
        )

        size_str = f"{item['bbox_width']}x{item['bbox_height']}"
        cv2.putText(
            canvas,
            f"Size: {size_str}",
            (x_start + 8, text_y + 34),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.38,
            (200, 200, 200),
            1,
            cv2.LINE_AA,
        )

        cv2.putText(
            canvas,
            f"Det: {item['method']}",
            (x_start + 8, text_y + 48),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.36,
            (170, 180, 195),
            1,
            cv2.LINE_AA,
        )

    cv2.imwrite(output_path, canvas)
    print(f"Saved visual grid image to '{output_path}'.")


def create_manual_review_excel(
    items: list[dict[str, Any]],
    output_path: str,
    sheet_title: str = "Review",
) -> None:
    """Create an Excel review worksheet with dropdown category validation."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    headers = ["sample_id", "probability", "review"]
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 4):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Setup Data Validation for dropdown
    cat_formula = '"' + ",".join(ALLOWED_CATEGORIES) + '"'
    dv = DataValidation(type="list", formula1=cat_formula, allow_blank=True)
    ws.add_data_validation(dv)

    row_start = 2
    for item in items:
        ws.append([item["sample_id"], round(item["probability"], 4), ""])
        current_row = ws.max_row

        c_id = ws.cell(row=current_row, column=1)
        c_p = ws.cell(row=current_row, column=2)
        c_r = ws.cell(row=current_row, column=3)

        c_id.alignment = align_left
        c_p.alignment = align_right
        c_p.number_format = "0.0000"
        c_r.alignment = align_center

    row_end = ws.max_row
    if row_end >= row_start:
        dv.add(f"C{row_start}:C{row_end}")

    # Set column widths
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 24

    wb.save(output_path)
    print(f"Saved Excel review sheet to '{output_path}'.")


def main() -> None:
    """Execute Phase 4D.1B manual contamination verification asset generation."""
    print("\n=======================================================")
    print("  PHASE 4D.1B MANUAL CONTAMINATION VERIFICATION")
    print("=======================================================")

    os.makedirs("audit", exist_ok=True)

    # Initialize perception components
    settings = get_settings()
    detector = CharacterDetector()
    roi_extractor = MapROIExtractor(settings=settings)
    classifier = VisualCandidateClassifier(model_path="models/entity_classifier.pkl", threshold=0.50)

    # Collect candidates across exploration and combat domains
    exp_folder = os.path.join("dataset", "exploration")
    com_folder = os.path.join("dataset", "combat")
    print("\nCollecting runtime candidates...")
    candidates = collect_runtime_candidates(exp_folder, com_folder, classifier, detector, roi_extractor)
    print(f"Total runtime candidates collected: {len(candidates)}")

    # Sort Top 100 Positives (P >= 0.50 descending)
    positives = [c for c in candidates if c["probability"] >= 0.50]
    positives.sort(key=lambda x: x["probability"], reverse=True)
    top100_positives = positives[:100]

    # Sort Borderline Candidates (0.40 <= P <= 0.60 ascending)
    borderline = [c for c in candidates if 0.40 <= c["probability"] <= 0.60]
    borderline.sort(key=lambda x: x["probability"])

    print(f"Top 100 accepted positives: {len(top100_positives)} items.")
    print(f"Borderline candidates (0.40 <= P <= 0.60): {len(borderline)} items.")

    # Task 1: Generate Visual Review Grids
    grid_top100_path = os.path.join("audit", "top100_positive_grid.png")
    render_image_grid(
        items=top100_positives,
        title="Phase 4D.1B — Top 100 Accepted Positives (P >= 0.50)",
        output_path=grid_top100_path,
        cols=10,
    )

    grid_borderline_path = os.path.join("audit", "borderline_review_grid.png")
    render_image_grid(
        items=borderline,
        title="Phase 4D.1B — Borderline Predictions Audit (0.40 <= P <= 0.60)",
        output_path=grid_borderline_path,
        cols=8,
    )

    # Task 2: Create audit/manual_positive_review.xlsx
    excel_pos_path = os.path.join("audit", "manual_positive_review.xlsx")
    create_manual_review_excel(
        items=top100_positives,
        output_path=excel_pos_path,
        sheet_title="Top 100 Positives",
    )

    # Task 3: Create audit/manual_borderline_review.xlsx
    excel_borderline_path = os.path.join("audit", "manual_borderline_review.xlsx")
    create_manual_review_excel(
        items=borderline,
        output_path=excel_borderline_path,
        sheet_title="Borderline Candidates",
    )

    print("\n=======================================================")
    print("  PHASE 4D.1B ASSETS SUCCESSFULLY GENERATED!")
    print("=======================================================")
    print("Files created:")
    print(f" - {grid_top100_path}")
    print(f" - {grid_borderline_path}")
    print(f" - {excel_pos_path}")
    print(f" - {excel_borderline_path}")
    print("=======================================================")


if __name__ == "__main__":
    main()
