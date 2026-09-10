"""Phase 4C.5 — Stress Candidate Extraction and Labeling Sheet Generator.

Extracts candidate crops from dataset/exploration and dataset/combat screenshots into audit/stress_candidates,
generates audit/stress_inventory.md, audit/manual_labels_stress.xlsx, and audit/stress_labeling_guide.md.

# ruff: noqa: N806, N803
"""

import glob
import json
import os
import random
from typing import Any

import cv2
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from dta.config.settings import get_settings
from dta.detectors.character_detector import CharacterDetector
from dta.vision.map_roi_extractor import MapROIExtractor

VALID_LABELS = [
    "player",
    "monster",
    "npc",
    "flower",
    "wall",
    "tree",
    "ground_tile",
    "movement_cell",
    "decoration",
    "ui_fragment",
    "box",
    "unknown",
]


def extract_candidates_from_folders(
    input_folders: list[str],
    output_dir: str,
) -> list[dict[str, Any]]:
    """Extract candidate crops from screenshot folders using CharacterDetector and save sidecars."""
    settings = get_settings()
    detector = CharacterDetector()
    roi_extractor = MapROIExtractor(settings=settings)

    os.makedirs(output_dir, exist_ok=True)
    extracted_records: list[dict[str, Any]] = []
    candidate_counter = 1

    for folder in input_folders:
        folder_name = os.path.basename(folder.rstrip("/\\"))
        png_files = sorted(glob.glob(os.path.join(folder, "*.png")))

        for fpath in png_files:
            fn = os.path.basename(fpath)
            frame_id = os.path.splitext(fn)[0]
            frame = cv2.imread(fpath)
            if frame is None or frame.size == 0:
                continue

            roi_frame = roi_extractor.apply_roi_mask(frame)
            trace = detector.detect_characters_with_trace(roi_frame)
            raw_characters = trace.get("raw_characters", [])

            for det in raw_characters:
                bbox = det.bbox
                x, y, w, h = max(0, bbox.x), max(0, bbox.y), bbox.w, bbox.h
                img_h, img_w = frame.shape[:2]
                x2, y2 = min(img_w, x + w), min(img_h, y + h)

                crop = frame[y:y2, x:x2]
                if crop.size == 0:
                    continue

                cid_str = f"{candidate_counter:04d}"
                img_filename = f"candidate_{cid_str}_{det.method}.png"
                json_filename = f"candidate_{cid_str}_{det.method}.json"

                crop_out_path = os.path.join(output_dir, img_filename)
                json_out_path = os.path.join(output_dir, json_filename)

                cv2.imwrite(crop_out_path, crop)

                sidecar_data = {
                    "candidate_id": candidate_counter,
                    "filename": img_filename,
                    "method": det.method,
                    "bbox": [x, y, w, h],
                    "area": float(det.area),
                    "confidence": float(det.confidence),
                    "frame_id": frame_id,
                    "source_folder": folder_name,
                    "diagnostics": det.diagnostics,
                }
                with open(json_out_path, "w", encoding="utf-8") as jf:
                    json.dump(sidecar_data, jf, indent=2)

                extracted_records.append({
                    "candidate_id": candidate_counter,
                    "cid_str": cid_str,
                    "filename": img_filename,
                    "method": det.method,
                    "frame_id": frame_id,
                    "source_folder": folder_name,
                    "crop_path": crop_out_path,
                    "sidecar_path": json_out_path,
                })

                candidate_counter += 1

    return extracted_records


def generate_inventory_report(
    records: list[dict[str, Any]],
    input_folders: list[str],
    output_path: str,
) -> None:
    """Generate audit/stress_inventory.md candidate inventory report."""
    total_count = len(records)
    method_counts = {"contour": 0, "hsv": 0, "combat_base": 0}
    folder_counts: dict[str, int] = {}
    unique_frames = set()

    for r in records:
        m = r["method"]
        method_counts[m] = method_counts.get(m, 0) + 1
        sf = r["source_folder"]
        folder_counts[sf] = folder_counts.get(sf, 0) + 1
        unique_frames.add(r["frame_id"])

    md = "# Phase 4C.5 — Stress Candidate Inventory Report\n\n"
    md += "## Executive Summary\n\n"
    md += f"- **Total Extracted Stress Candidate Crops**: **{total_count}**\n"
    md += f"- **Unique Source Screenshots Audited**: **{len(unique_frames)}**\n"
    md += "- **Target Output Location**: `audit/stress_candidates/`\n\n"

    md += "## 1. Candidate Breakdown by Detector Generator Method\n\n"
    md += "| Detector Method | Candidate Count | Share (%) |\n"
    md += "| :--- | :---: | :---: |\n"
    for m, count in method_counts.items():
        pct = (count / float(max(1, total_count))) * 100.0
        md += f"| `{m}` | **{count}** | {pct:.1f}% |\n"

    md += "\n## 2. Candidate Breakdown by Source Folder\n\n"
    md += "| Source Directory | Candidate Crops Extracted |\n"
    md += "| :--- | :---: |\n"
    for sf, count in folder_counts.items():
        md += f"| `dataset/{sf}` | **{count}** |\n"

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(md)

    print(f"Generated candidate inventory report at '{output_path}'.")


def create_stratified_labeling_sheet(
    records: list[dict[str, Any]],
    output_excel_path: str,
    sample_size: int = 150,
) -> list[dict[str, Any]]:
    """Create audit/manual_labels_stress.xlsx with stratified sampling balanced across detector methods."""
    random.seed(42)

    # Group records by method
    by_method: dict[str, list[dict[str, Any]]] = {
        "contour": [],
        "hsv": [],
        "combat_base": [],
    }
    for r in records:
        m = r["method"]
        if m in by_method:
            by_method[m].append(r)

    # Stratified sampling target per method (~50 each for N=150)
    per_method_target = sample_size // 3
    sampled_records: list[dict[str, Any]] = []

    for _m, item_list in by_method.items():
        n_available = len(item_list)
        n_take = min(n_available, per_method_target)
        shuffled = list(item_list)
        random.shuffle(shuffled)
        sampled_records.extend(shuffled[:n_take])

    # Sort sampled records by numerical candidate_id
    sampled_records.sort(key=lambda r: r["candidate_id"])

    # Create Excel Workbook
    wb = openpyxl.Workbook()

    # 1. Main Sheet
    ws = wb.active
    ws.title = "StressLabels"

    # 2. Metadata Sheet for validation dropdown
    ws_meta = wb.create_sheet(title="Metadata")
    for idx, lbl in enumerate(VALID_LABELS, start=1):
        ws_meta.cell(row=idx, column=1, value=lbl)
    ws_meta.freeze_panes = "A2"

    # Header
    headers = ["candidate_id", "filename", "method", "source_folder", "accepted", "label", "notes"]
    ws.append(headers)

    for r in sampled_records:
        ws.append([
            r["candidate_id"],
            r["filename"],
            r["method"],
            r["source_folder"],
            "accepted",
            "",  # empty for manual annotation
            "",
        ])

    num_rows = len(sampled_records) + 1

    # Data Validation Dropdown for "label" (Column F)
    dv = DataValidation(
        type="list",
        formula1=f"=Metadata!$A$1:$A${len(VALID_LABELS)}",
        allow_blank=True,
    )
    dv.error = "Select a valid label from dropdown."
    dv.errorTitle = "Invalid Label"
    dv.prompt = "Select label"
    dv.promptTitle = "Label Choice"

    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(2, num_rows)}")

    # Format Table
    table_ref = f"A1:G{max(2, num_rows)}"
    tab = Table(displayName="StressLabelsTable", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)
    ws.freeze_panes = "A2"

    # Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_excel_path)
    print(f"Created stratified labeling sheet '{output_excel_path}' with {len(sampled_records)} sampled candidates.")

    return sampled_records


def generate_labeling_guide(output_path: str) -> None:
    """Generate audit/stress_labeling_guide.md defining ground truth labeling taxonomy."""
    guide_content = r"""# Phase 4C.5 — Stress Candidate Labeling Guide

This guide defines the ground-truth annotation taxonomy for labeling candidate crops in `audit/manual_labels_stress.xlsx`.

---

## 1. Ground Truth Binary Mapping

All labels fall strictly into two target ground-truth classes:

$$\\text{Ground Truth} = \\begin{cases} 1 & \\text{if True Entity (TP: player, monster, npc)} \\\\ 0 & \\text{if Scenery / Prop / Noise (FP)} \\end{cases}$$

---

## 2. Category Taxonomy & Definitions

### True Positives ($\text{Label} = 1$)

- **`player`**: Player character sprite (full sprite or head/body crop).
- **`monster`**: Aggressive or neutral mob/monster sprite on exploration map or combat grid.
- **`npc`**: Non-player character entity, shop vendor, or interactive character.

### False Positives ($\text{Label} = 0$)

- **`flower`**: Map flowers, grass bushes, crop plants, or environmental flora.
- **`wall`**: Wall fragments, stone borders, building edges, or fence posts.
- **`tree`**: Tree trunks, leaves, canopy branches, or forest vegetation.
- **`ground_tile`**: Floor tiles, dirt paths, cobblestone textures, or water ripples.
- **`movement_cell`**: Green PM movement grid tiles or combat range highlights.
- **`decoration`**: Statues, lamps, flags, signs, rocks, or static map ornaments.
- **`ui_fragment`**: HUD banners, action bar icons, minimap edges, or text overlays.
- **`box`**: Wooden crates, barrels, chests, or interactive containers.
- **`unknown`**: Ambiguous scenery artifacts or noise fragments.

---

## 3. Labeling Procedure in Excel

1. Open [`audit/manual_labels_stress.xlsx`](file:///c:/Users/Andres/Desktop/bot/audit/manual_labels_stress.xlsx) in Excel / LibreOffice.
2. For each row, inspect the crop image stored at `audit/stress_candidates/<filename>`.
3. Select the precise category from the **`label`** dropdown in Column F.
4. Save the completed workbook.
"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(guide_content)

    print(f"Generated labeling guide at '{output_path}'.")


def main() -> None:
    """Execute candidate extraction, inventory report, labeling sheet, and guide creation."""
    input_folders = [
        os.path.join("dataset", "exploration"),
        os.path.join("dataset", "combat"),
    ]
    output_dir = os.path.join("audit", "stress_candidates")

    print("\n=======================================================")
    print("  PHASE 4C.5 STRESS CANDIDATE EXTRACTION")
    print("=======================================================")

    # 1. Extract candidates
    records = extract_candidates_from_folders(input_folders, output_dir)
    print(f"Extracted total of {len(records)} candidate crops to '{output_dir}'.")

    # 2. Generate inventory report
    inventory_path = os.path.join("audit", "stress_inventory.md")
    generate_inventory_report(records, input_folders, inventory_path)

    # 3. Create stratified labeling sheet (N=150)
    excel_path = os.path.join("audit", "manual_labels_stress.xlsx")
    create_stratified_labeling_sheet(records, excel_path, sample_size=150)

    # 4. Generate labeling guide
    guide_path = os.path.join("audit", "stress_labeling_guide.md")
    generate_labeling_guide(guide_path)

    print("\nPhase 4C.5 Candidate Dataset Preparation Complete!")


if __name__ == "__main__":
    main()
